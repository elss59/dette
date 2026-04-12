#!/usr/bin/env python3
import json
from pathlib import Path

INPUT = Path("data/dette.xlsx")
OUTPUT = Path("src/data/dette.generated.json")
SHEET_NAME = "Dette"

HEADER_MAP = {
    "Produit": "produit",
    "Parcours": "parcours",
    "Impact expérientiel": "impactExperientiel",
    "Date dette": "dateDette",
    "Lien Figma cible": "figmaCible",
    'Explain "Version cible"': "descriptionCible",
    "Lien Figma inter": "figmaIntermediaire",
    'Explain "Version intermédiaire"': "descriptionIntermediaire",
    "Lien vidéo prod": "videoProd",
    "Explain écart prod": "descriptionProd",
    "Raison de l'écart": "raisonEcart",
    "Impact user": "impactUtilisateur",
    "Commentaire raison de l'écart": "commentaire",
    "Accès Vision": "accesVision",
    "Titre Vision": "titreVision",
}

TEMPLATE = {value: "" for value in HEADER_MAP.values()}

if not INPUT.exists():
    print(f"[dette:xlsx] Fichier introuvable: {INPUT}")
    raise SystemExit(0)

try:
    from openpyxl import load_workbook
except ImportError:
    print("[dette:xlsx] openpyxl non disponible. Installez-le pour activer la conversion Excel.")
    raise SystemExit(0)

workbook = load_workbook(INPUT, data_only=True)
if SHEET_NAME not in workbook.sheetnames:
    raise ValueError(f'[dette:xlsx] Feuille "{SHEET_NAME}" introuvable.')

sheet = workbook[SHEET_NAME]
rows = list(sheet.iter_rows(values_only=True))
if not rows:
    OUTPUT.write_text("[]\n", encoding="utf-8")
    print(f"[dette:xlsx] 0 ligne convertie vers {OUTPUT}")
    raise SystemExit(0)

headers = [str(h).strip() if h is not None else "" for h in rows[0]]
entries = []
for row in rows[1:]:
    if row is None:
        continue
    row_dict = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
    if not any((value not in (None, "")) for value in row_dict.values()):
        continue

    mapped = dict(TEMPLATE)
    for excel_header, target_key in HEADER_MAP.items():
        value = row_dict.get(excel_header, "")
        mapped[target_key] = "" if value is None else str(value)
    entries.append(mapped)

OUTPUT.write_text(json.dumps(entries, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"[dette:xlsx] {len(entries)} ligne(s) convertie(s) vers {OUTPUT}")
