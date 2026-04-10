#!/usr/bin/env python3
import json
from pathlib import Path
from datetime import datetime, date

SOURCE = Path('/data/dette.xlsx')
OUTPUT = Path('src/data/detteDataset.json')
SHEET_NAME = 'Dette'

FIELDS = [
    'produit',
    'parcours',
    'impactExperientiel',
    'dateDette',
    'figmaCible',
    'descriptionCible',
    'figmaIntermediaire',
    'descriptionIntermediaire',
    'videoProd',
    'descriptionProd',
    'raisonEcart',
    'impactUtilisateur',
    'commentaire',
    'accesVision',
    'titreVision',
]


def normalize(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()

rows = []

if SOURCE.exists():
    from openpyxl import load_workbook

    wb = load_workbook(SOURCE, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Sheet '{SHEET_NAME}' introuvable dans {SOURCE}")

    ws = wb[SHEET_NAME]
    header = [normalize(c.value) for c in ws[1]]
    idx = {name: i for i, name in enumerate(header)}

    for expected in FIELDS:
        if expected not in idx:
            raise SystemExit(f"Colonne manquante dans la feuille '{SHEET_NAME}': {expected}")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None:
            continue
        obj = {field: normalize(row[idx[field]]) for field in FIELDS}
        if any(obj.values()):
            rows.append(obj)
else:
    print(f"[WARN] Fichier source absent: {SOURCE}. Génération d'un dataset vide.")

OUTPUT.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(f"JSON généré: {OUTPUT} ({len(rows)} lignes)")
